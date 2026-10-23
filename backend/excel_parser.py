"""
Parser de archivos Excel para carga de datos — INVESAKK SAS
Acepta formato SQL directo y formato clásico.
"""
from openpyxl import load_workbook
from datetime import date, datetime
import re, os

CONDUCTORES_PLACAS = {
    # Placas actuales
    'WGX-062': 'ROA',      'WGX062': 'ROA',
    'WGX-060': 'MARTINEZ', 'WGX060': 'MARTINEZ',
    'TDU-499': 'MARIN',    'TDU499': 'MARIN',
    'WGV-233': 'MENDOZA',  'WGV233': 'MENDOZA',
    'WGX-061': 'NIÑO',     'WGX061': 'NIÑO',
    'WGD-149': 'RICO',     'WGD149': 'RICO',
}
CONDUCTORES_KEYS = list(CONDUCTORES_PLACAS.values())


def nc(s):
    """Normaliza nombre de columna: minúsculas sin espacios ni guiones."""
    return re.sub(r'[\s_\-]+', '', str(s or '').strip().lower())


def col(row, *names):
    keys = list(row.keys())
    norm = {nc(k): k for k in keys}
    for name in names:
        n = nc(name)
        # coincidencia exacta normalizada
        if n in norm:
            return row[norm[n]]
        # coincidencia parcial (primeros 6 chars)
        match = next((orig for orig_n, orig in norm.items() if orig_n.startswith(n[:6])), None)
        if match:
            return row[match]
    return None


def fmt_excel_date(v):
    if not v:
        return None
    if isinstance(v, (date, datetime)):
        return v.strftime('%Y-%m-%d')
    s = str(v).strip()
    # YYYY-MM-DD
    if re.match(r'^\d{4}-\d{2}-\d{2}', s):
        return s[:10]
    # serial Excel
    try:
        n = float(s)
        if n > 40000:
            from datetime import timedelta
            base = date(1899, 12, 30)
            return (base + timedelta(days=int(n))).strftime('%Y-%m-%d')
    except (ValueError, TypeError):
        pass
    return None


def detect_cond_key(cond_raw, placa_raw=''):
    """Busca la clave del conductor por nombre o placa."""
    c = str(cond_raw or '').strip().upper()
    p = str(placa_raw or '').strip().upper()

    # Placa exacta
    if p in CONDUCTORES_PLACAS:
        return CONDUCTORES_PLACAS[p]

    # Clave exacta
    if c in CONDUCTORES_KEYS:
        return c

    # Nombre contiene la clave
    for k in CONDUCTORES_KEYS:
        if k in c:
            return k

    return None


def tipo_es_mt(tipo):
    return bool(re.match(r'^NT', str(tipo or '').strip(), re.IGNORECASE))


# ── Festivos Colombia (actualizables anualmente) ───────────────
_FESTIVOS_CO = {
    date(a, m, d)
    for d, m, a in [
        # 2025
        (1,1,2025),(6,1,2025),(24,3,2025),(17,4,2025),(18,4,2025),
        (1,5,2025),(2,6,2025),(23,6,2025),(30,6,2025),(20,7,2025),
        (7,8,2025),(18,8,2025),(13,10,2025),(3,11,2025),(17,11,2025),
        (8,12,2025),(25,12,2025),
        # 2026
        (1,1,2026),(12,1,2026),(23,3,2026),(2,4,2026),(3,4,2026),
        (1,5,2026),(18,5,2026),(8,6,2026),(15,6,2026),(29,6,2026),
        (20,7,2026),(7,8,2026),(17,8,2026),(12,10,2026),(2,11,2026),
        (16,11,2026),(8,12,2026),(25,12,2026),
        # 2027 — Easter: Mar 28
        (1,1,2027),(11,1,2027),(22,3,2027),(25,3,2027),(26,3,2027),
        (1,5,2027),(10,5,2027),(31,5,2027),(7,6,2027),(5,7,2027),
        (20,7,2027),(7,8,2027),(16,8,2027),(18,10,2027),(1,11,2027),
        (15,11,2027),(8,12,2027),(25,12,2027),
    ]
}


def dias_habiles(fecha_inicio_str, fecha_fin_str):
    """
    Días hábiles entre dos fechas (lun–sáb, sin festivos colombianos).
    Positivo = entrega tarde, negativo = entrega anticipada.
    """
    try:
        from datetime import timedelta
        fi = date.fromisoformat(fecha_inicio_str[:10])
        ff = date.fromisoformat(fecha_fin_str[:10])
        if fi == ff:
            return 0
        sign = 1 if ff > fi else -1
        start, end = (fi, ff) if sign == 1 else (ff, fi)
        habiles = 0
        cur = start + timedelta(days=1)
        while cur <= end:
            if cur.weekday() != 6 and cur not in _FESTIVOS_CO:  # 6=domingo
                habiles += 1
            cur += timedelta(days=1)
        return sign * habiles
    except Exception:
        return None


def parse_excel(filepath, tipo):
    wb = load_workbook(filepath, read_only=True, data_only=True)

    # Intentar hoja por nombre, si no la primera
    sheet_map = {
        'facturas':    ['Facturas', 'Sheet1', 'Hoja1', 'v_ventas_vendedores_ppto'],
        'conductores': ['Conductores', 'Sheet2', 'Hoja2'],
        'gps':         ['GPS', 'Sheet3', 'Hoja3'],
        'bodegas':     ['Bodegas', 'Sheet4', 'Hoja4'],
    }
    ws = None
    for name in sheet_map.get(tipo, []):
        if name in wb.sheetnames:
            candidate = wb[name]
            if candidate.max_row and candidate.max_row >= 2:
                ws = candidate
                break
    if ws is None:
        ws = wb.active

    rows = list(ws.iter_rows(values_only=True))
    # Si la hoja seleccionada está vacía, buscar la primera con datos
    if len(rows) < 2:
        for name in wb.sheetnames:
            candidate = wb[name]
            crows = list(candidate.iter_rows(values_only=True))
            if len(crows) >= 2:
                rows = crows
                break
    if len(rows) < 2:
        return [], "El archivo no tiene datos suficientes"

    # Buscar fila de encabezado (puede haber filas de título arriba)
    header_row = None
    for i, row in enumerate(rows[:5]):
        vals = [str(v or '').strip() for v in row]
        non_empty = [v for v in vals if v]
        if len(non_empty) >= 2:
            header_row = i
            break

    if header_row is None:
        return [], "No se encontró fila de encabezados"

    headers = [str(v or '') for v in rows[header_row]]
    data_rows = []

    for row in rows[header_row + 1:]:
        if all(v is None or str(v).strip() == '' for v in row):
            continue
        data_rows.append(dict(zip(headers, row)))

    if not data_rows:
        return [], "No hay filas de datos"

    if tipo == 'facturas':
        return _parse_facturas(data_rows)
    elif tipo == 'conductores':
        return _parse_conductores(data_rows)
    elif tipo == 'gps':
        return _parse_gps(data_rows)
    elif tipo == 'bodegas':
        return _parse_bodegas(data_rows)
    else:
        return [], f"Tipo desconocido: {tipo}"


def _parse_facturas(data_rows):
    results = []
    for r in data_rows:
        # Conductor: columna 'nombre' (SQL) o 'conductor' (clásico)
        cond_raw  = col(r, 'nombre', 'conductor') or ''
        placa_raw = col(r, 'placa') or ''
        cond_key  = detect_cond_key(cond_raw, placa_raw) or 'SIN_ASIGNAR'

        tipo     = str(col(r, 'tipo', 'tipodoc') or '').strip().upper()
        tip_desc = str(col(r, 'descripcion') or '').strip()
        es_mt    = tipo_es_mt(tipo) or bool(re.match(r'^si$', str(col(r, 'esnt', 'nt') or ''), re.I))

        # Fechas
        fec_fact    = fmt_excel_date(col(r, 'fecha', 'fecfactura', 'fechafactura', 'fecha_factura', 'fecha_hora', 'fechahora'))
        fec_promesa = fmt_excel_date(col(r, 'fecpromesaentrega', 'fec_promesa_entrega', 'fechapromesa'))
        fec_entr    = fmt_excel_date(col(r, 'fecentreganegocio', 'fec_entrega_logistica', 'fechaentrega', 'fecha_entrega'))

        # Diferencia días — recalcular en días hábiles si hay ambas fechas
        dif_raw  = col(r, 'diferenciafechadias', 'diferencia_fecha_dias', 'difdias')
        dif_dias = float(dif_raw) if dif_raw not in (None, '', 'None') else None
        if fec_promesa and fec_entr:
            dif_dias = dias_habiles(fec_promesa, fec_entr)
        dif_h_r  = col(r, 'diferenciafechahoras', 'diferencia_fecha_horas')
        dif_horas= float(dif_h_r) if dif_h_r not in (None, '', 'None') else None

        # Valor / peso
        valor = col(r, 'total_trans', 'valortotal', 'valor_total', 'valor')
        peso  = col(r, 'peso_total', 'pesototal', 'peso')

        results.append({
            'cond_key':  cond_key,
            'num':       str(col(r, 'numero', 'numdoc', 'factura', 'num') or '—'),
            'tipo':      tipo,
            'tip_desc':  tip_desc,
            'cliente':   str(col(r, 'nombres', 'cliente') or ''),
            'ciudad':    str(col(r, 'ciudad') or ''),
            'dpto':      str(col(r, 'dpto', 'departamento') or ''),
            'bodega':    str(col(r, 'bodega', 'sede') or ''),
            'placa':     str(placa_raw),
            'valor':     float(valor) if valor not in (None, '', 'None') else None,
            'peso':      float(peso)  if peso  not in (None, '', 'None') else None,
            'fec_fact':  fec_fact,
            'fec_promesa': fec_promesa,
            'fec_entr':  fec_entr,
            'dif_dias':  dif_dias,
            'dif_horas': dif_horas,
            'es_mt':     es_mt,
        })
    return results, None


def _parse_conductores(data_rows):
    results = []
    for r in data_rows:
        clave = str(col(r, 'clave') or '').strip().upper()
        if not clave:
            continue
        results.append({
            'clave':    clave,
            'nombre':   str(col(r, 'nombre') or clave),
            'short':    str(col(r, 'short') or clave[:5]),
            'placa':    str(col(r, 'placa') or ''),
            'pedidos':  int(col(r, 'pedidos') or 0),
            'a_tiempo': int(col(r, 'atiempo', 'a_tiempo') or 0),
            'h24':      int(col(r, 'h24') or 0),
            'd2':       int(col(r, 'd2') or 0),
            'bono_c':   float(col(r, 'bonoc', 'bono_c') or 0),
            'bono_a':   float(col(r, 'bonoa', 'bono_a') or 0),
            'meta_c':   float(col(r, 'metac', 'meta_c') or 0),
            'meta_a':   float(col(r, 'metaa', 'meta_a') or 0),
            'meta_h24': int(col(r, 'metah24') or 0),
        })
    return results, None


def _parse_gps(data_rows):
    results = []
    for r in data_rows:
        placa = str(col(r, 'placa') or '').strip()
        if not placa:
            continue
        results.append({
            'placa':     placa,
            'conductor': str(col(r, 'conductor') or ''),
            'vel':       int(col(r, 'excesvel', 'exceso', 'exvel') or 0),
            'acc':       int(col(r, 'acelbrusca', 'aceleracion') or 0),
            'brk':       int(col(r, 'frenadobrusco', 'frenado') or 0),
        })
    return results, None


def _parse_bodegas(data_rows):
    results = []
    for r in data_rows:
        sede = str(col(r, 'sede') or '').strip()
        if not sede:
            continue
        results.append({
            'sede': sede,
            'ped':  int(col(r, 'pedidos') or 0),
            'nt':   int(col(r, 'nt') or 0),
        })
    return results, None
