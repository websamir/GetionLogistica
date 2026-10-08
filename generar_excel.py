"""
Genera el archivo Excel de plantilla para carga en Gestión Logística INVESAKK SAS
Formato: exportación directa SQL
"""
import openpyxl
from openpyxl.styles import (
    Font, PatternFill, Alignment, Border, Side, GradientFill
)
from openpyxl.utils import get_column_letter
from datetime import date, timedelta
import os

OUT = r"D:\Users\lgamarra\Desktop\hub\GestionLogistica\Plantilla_Facturas_Logistica.xlsx"

wb = openpyxl.Workbook()

# ── Colores corporativos ──────────────────────────────────────
AZUL_OSCURO = "1D4ED8"
AZUL_MED    = "3B82F6"
AZUL_CLARO  = "DBEAFE"
GRIS_HEADER = "F1F5F9"
VERDE       = "059669"
NARANJA     = "D97706"
ROJO        = "DC2626"
MORADO      = "7C3AED"
BLANCO      = "FFFFFF"
GRIS_BORDE  = "CBD5E1"
AMARILLO_S  = "FEF9C3"
VERDE_S     = "DCFCE7"
ROJO_S      = "FEE2E2"
MORADO_S    = "EDE9FE"
NEUTRO_S    = "F1F5F9"

def borde_fino():
    s = Side(style='thin', color=GRIS_BORDE)
    return Border(left=s, right=s, top=s, bottom=s)

def fill(hex_color):
    return PatternFill("solid", fgColor=hex_color)

# ════════════════════════════════════════════════════════════════
# HOJA 1: FACTURAS  (formato SQL)
# ════════════════════════════════════════════════════════════════
ws = wb.active
ws.title = "Facturas"
ws.sheet_view.showGridLines = False

# Fila 1 — Título
ws.merge_cells("A1:O1")
ws["A1"] = "INVESAKK SAS · Plantilla de Facturas / Pedidos — Gestión Logística"
ws["A1"].font      = Font(name="Calibri", bold=True, size=13, color=BLANCO)
ws["A1"].fill      = fill(AZUL_OSCURO)
ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
ws.row_dimensions[1].height = 28

# Fila 2 — Sub
ws.merge_cells("A2:O2")
ws["A2"] = "Exportación SQL · sube este archivo en el panel Admin → pestaña Facturas/Pedidos"
ws["A2"].font      = Font(name="Calibri", italic=True, size=10, color="94A3B8")
ws["A2"].fill      = fill("1E3A8A")
ws["A2"].alignment = Alignment(horizontal="center", vertical="center")
ws.row_dimensions[2].height = 18

ws.row_dimensions[3].height = 6   # espaciado

# Columnas: nombre exacto que acepta la plataforma (SQL)
COLS = [
    # (encabezado,            ancho, color_fondo,  tooltip)
    ("numero",                18,    AZUL_CLARO,   "Número de documento (factura, nota, etc.)"),
    ("nombre",                24,    AZUL_CLARO,   "Nombre completo del conductor (ej. MARIN (TR) JOSE)"),
    ("tipo",                  10,    NEUTRO_S,     "Tipo: NT20 | NT10 | NT6 | FEPA | G1AG"),
    ("descripcion",           26,    NEUTRO_S,     "Descripción del tipo de documento"),
    ("nombres",               30,    GRIS_HEADER,  "Nombre del cliente / tercero"),
    ("ciudad",                14,    GRIS_HEADER,  "Ciudad de entrega"),
    ("dpto",                  12,    GRIS_HEADER,  "Departamento"),
    ("placa",                 10,    NEUTRO_S,     "Placa del vehículo (alternativa para identificar conductor)"),
    ("total_trans",           14,    VERDE_S,      "Valor total de la transacción (COP)"),
    ("peso_total",            12,    VERDE_S,      "Peso total en kg"),
    ("fec_Factura",           14,    AMARILLO_S,   "Fecha de factura · formato YYYY-MM-DD"),
    ("fec_promesa_entrega",   22,    AMARILLO_S,   "Fecha prometida de entrega · YYYY-MM-DD"),
    ("fec_entrega_logistica", 22,    AMARILLO_S,   "Fecha real de entrega logística · vacía = pendiente"),
    ("diferencia_fecha_dias", 22,    ROJO_S,       "Diferencia en días (calculada por SQL) · opcional"),
    ("diferencia_fecha_horas",22,    ROJO_S,       "Diferencia en horas · opcional"),
]

# Encabezados (fila 4)
for ci, (col_name, ancho, bg, _) in enumerate(COLS, 1):
    cell = ws.cell(row=4, column=ci, value=col_name)
    cell.font      = Font(name="Calibri", bold=True, size=10,
                          color=AZUL_OSCURO if bg != AZUL_CLARO else BLANCO if bg == AZUL_OSCURO else AZUL_OSCURO)
    cell.fill      = fill(AZUL_MED if bg == AZUL_CLARO else bg)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border    = borde_fino()
    ws.column_dimensions[get_column_letter(ci)].width = ancho
ws.row_dimensions[4].height = 36

# Encabezados con color azul para columnas clave
for ci in [1, 2]:
    ws.cell(row=4, column=ci).fill = fill(AZUL_MED)
    ws.cell(row=4, column=ci).font = Font(name="Calibri", bold=True, size=10, color=BLANCO)

# ── Datos de ejemplo (23 filas MARIN) ────────────────────────────
CONDUCTORES = {
    "ROA":      ("ROA (TR) JULIO ARMANDO", "WGX062"),
    "MARTINEZ": ("MARTINEZ (TR) LUIS",      "WGX060"),
    "MARIN":    ("MARIN (TR) JOSE",         "TDU-499"),
    "MENDOZA":  ("MENDOZA (TR) LUIS",       "WGX061"),
    "NIÑO":     ("NIÑO (TR) ADEMIR",        "WGD149"),
}

# Tipos y descripciones
TIPOS = [
    ("NT20", "NT Licitaciones"),
    ("NT10", "NT Caribe Verde"),
    ("FEPA", "Factura Electrónica"),
    ("G1AG", "Fac Convenios"),
    ("NT6",  "NT Bodemayor"),
]

CIUDADES = [
    ("Barranquilla", "Atlántico"),
    ("Malambo",      "Atlántico"),
    ("Soledad",      "Atlántico"),
    ("Puerto Colombia", "Atlántico"),
    ("Galapa",       "Atlántico"),
]

CLIENTES = [
    "FERRETERÍA EL PROGRESO",
    "MATERIALES DE CONSTRUCCIÓN ABC",
    "INVERSIONES CARIBE SAS",
    "ALMACENES BLOCK",
    "CONSTRUCTORA DEL NORTE",
    "GRUPO FERRETERO DEL CARIBE",
    "DISTRIBUIDORA EL MAESTRO",
    "CERÁMICA Y ACABADOS SAS",
    "FERROMUNDO LTDA",
    "PINTURAS DEL ATLÁNTICO",
]

import random
random.seed(42)

sample_rows = []
# 23 filas de ejemplo para MARIN, más 5 por cada otro conductor
for cond_key, (cond_nombre, placa) in CONDUCTORES.items():
    n_rows = 23 if cond_key == "MARIN" else 5
    for i in range(n_rows):
        tipo, tipdesc = random.choice(TIPOS)
        ciudad, dpto  = random.choice(CIUDADES)
        cliente        = random.choice(CLIENTES)
        val            = random.randint(80, 1500) * 1000
        peso           = round(random.uniform(50, 800), 1)

        # Fecha factura: dentro de septiembre 2026
        dia_fact  = random.randint(1, 25)
        fec_fact  = date(2026, 9, dia_fact)
        fec_prom  = fec_fact + timedelta(days=random.randint(0, 3))

        esMT = tipo.startswith("NT")
        if esMT:
            dif_dias  = None
            dif_horas = None
            fec_entr  = None
        else:
            dif_dias  = random.choice([0,0,1,2,2,3,4,5,5,7,10,15,25])
            dif_horas = dif_dias * 24 + random.randint(0, 23) if dif_dias is not None else None
            fec_entr  = fec_prom + timedelta(days=dif_dias) if dif_dias is not None else None

        num = f"FAC-{random.randint(10000,99999)}"

        sample_rows.append([
            num,
            cond_nombre,
            tipo,
            tipdesc,
            cliente,
            ciudad,
            dpto,
            placa,
            val,
            peso,
            fec_fact.strftime("%Y-%m-%d"),
            fec_prom.strftime("%Y-%m-%d"),
            fec_entr.strftime("%Y-%m-%d") if fec_entr else "",
            dif_dias if dif_dias is not None else "",
            dif_horas if dif_horas is not None else "",
        ])

# Colores de fondo por tipo
TIPO_FILL = {
    "NT20": NEUTRO_S, "NT10": NEUTRO_S, "NT6": NEUTRO_S,
    "FEPA": VERDE_S,  "G1AG": MORADO_S,
}

for ri, row in enumerate(sample_rows, 5):
    tipo_val = row[2]
    row_bg   = TIPO_FILL.get(tipo_val, BLANCO)
    alt_bg   = "F8FAFC" if ri % 2 == 0 else BLANCO

    for ci, val in enumerate(row, 1):
        cell = ws.cell(row=ri, column=ci, value=val)
        cell.font      = Font(name="Calibri", size=10)
        cell.border    = borde_fino()
        cell.alignment = Alignment(vertical="center")

        # Columnas de fecha: centradas
        if ci in [11, 12, 13]:
            cell.alignment = Alignment(horizontal="center", vertical="center")
        # Valor y peso: formato número
        if ci == 9 and val:
            cell.number_format = '$ #,##0'
        if ci == 10 and val:
            cell.number_format = '#,##0.0" kg"'
        if ci == 14 and val != "":
            cell.number_format = '0" días"'
        # Color de fondo alternado con color de tipo
        if tipo_val.startswith("NT"):
            bg = NEUTRO_S if ri % 2 == 0 else "F8FAFC"
        else:
            bg = VERDE_S if tipo_val == "FEPA" else MORADO_S if tipo_val == "G1AG" else alt_bg
        cell.fill = fill(bg)
    ws.row_dimensions[ri].height = 18

# Freeze panes y filtro
ws.freeze_panes = "A5"
ws.auto_filter.ref = f"A4:{get_column_letter(len(COLS))}4"

# ════════════════════════════════════════════════════════════════
# HOJA 2: CONDUCTORES
# ════════════════════════════════════════════════════════════════
ws2 = wb.create_sheet("Conductores")
ws2.sheet_view.showGridLines = False

ws2.merge_cells("A1:I1")
ws2["A1"] = "INVESAKK SAS · Datos de Conductores"
ws2["A1"].font = Font(name="Calibri", bold=True, size=13, color=BLANCO)
ws2["A1"].fill = fill(VERDE)
ws2["A1"].alignment = Alignment(horizontal="center", vertical="center")
ws2.row_dimensions[1].height = 28

COLS2 = [
    ("Clave",    10, "Clave de conductor: ROA | MARTINEZ | MARIN | MENDOZA | NIÑO"),
    ("Nombre",   28, "Nombre completo del conductor"),
    ("Placa",    12, "Placa del vehículo"),
    ("Pedidos",  10, "Total pedidos del período"),
    ("ATiempo",  10, "Pedidos entregados a tiempo (≤5 días)"),
    ("H24",      8,  "Pedidos entregados en ≤24 horas"),
    ("D2",       8,  "Pedidos en ≤2 días"),
    ("BonoC",    14, "Bono calculado conductor (COP)"),
    ("BonoA",    14, "Bono calculado ayudante (COP)"),
]
for ci, (name, ancho, _) in enumerate(COLS2, 1):
    c = ws2.cell(row=2, column=ci, value=name)
    c.font = Font(name="Calibri", bold=True, size=10, color=BLANCO)
    c.fill = fill(VERDE)
    c.alignment = Alignment(horizontal="center", vertical="center")
    c.border = borde_fino()
    ws2.column_dimensions[get_column_letter(ci)].width = ancho
ws2.row_dimensions[2].height = 30

data_cond = [
    ["ROA",      "Julio Armando Roa",  "WGX062", 170, 77,  0, 11, 264200,  132100],
    ["MARTINEZ", "Luis Martínez",      "WGX060", 147, 59,  0, 11, 201000,  100500],
    ["MARIN",    "José Marín",         "TDU-499",  70, 18,  1,  0,  35200,   17600],
    ["MENDOZA",  "Luis Mendoza",       "WGX061", 200, 14,  0,  1,-110600,  -55300],
    ["NIÑO",     "Ademir Niño",        "WGD149",  66, 11,  0,  3,  -1400,    -700],
]
COND_COLS = ["#1D4ED8","#7C3AED","#059669","#D97706","#DC2626"]
for ri, row in enumerate(data_cond, 3):
    for ci, val in enumerate(row, 1):
        c = ws2.cell(row=ri, column=ci, value=val)
        c.font = Font(name="Calibri", size=10, bold=(ci==1))
        c.border = borde_fino()
        c.alignment = Alignment(vertical="center")
        bg = VERDE_S if ri % 2 == 0 else BLANCO
        c.fill = fill(bg)
        if ci in [8,9]:
            c.number_format = '$ #,##0;[Red]$ -#,##0'
    ws2.row_dimensions[ri].height = 18

ws2.freeze_panes = "A3"

# ════════════════════════════════════════════════════════════════
# HOJA 3: GPS
# ════════════════════════════════════════════════════════════════
ws3 = wb.create_sheet("GPS")
ws3.sheet_view.showGridLines = False

ws3.merge_cells("A1:E1")
ws3["A1"] = "INVESAKK SAS · Eventos GPS / Comportamiento Vial"
ws3["A1"].font = Font(name="Calibri", bold=True, size=13, color=BLANCO)
ws3["A1"].fill = fill(NARANJA)
ws3["A1"].alignment = Alignment(horizontal="center", vertical="center")
ws3.row_dimensions[1].height = 28

COLS3 = [("Placa",12),("Conductor",26),("ExcesoVel",14),("AcelBrusca",14),("FrenadoBrusco",16)]
for ci,(name,ancho) in enumerate(COLS3,1):
    c = ws3.cell(row=2, column=ci, value=name)
    c.font = Font(name="Calibri", bold=True, size=10, color=BLANCO)
    c.fill = fill(NARANJA)
    c.alignment = Alignment(horizontal="center",vertical="center")
    c.border = borde_fino()
    ws3.column_dimensions[get_column_letter(ci)].width = ancho
ws3.row_dimensions[2].height = 30

gps_data = [
    ["WGX062","Julio Armando Roa",2,1,3],
    ["WGX060","Luis Martínez",    1,0,2],
    ["TDU-499","José Marín",      0,0,1],
    ["WGX061","Luis Mendoza",     4,2,5],
    ["WGD149","Ademir Niño",      1,1,2],
]
for ri, row in enumerate(gps_data, 3):
    for ci, val in enumerate(row, 1):
        c = ws3.cell(row=ri, column=ci, value=val)
        c.font = Font(name="Calibri", size=10)
        c.border = borde_fino()
        c.alignment = Alignment(vertical="center")
        c.fill = fill(AMARILLO_S if ri%2==0 else BLANCO)
    ws3.row_dimensions[ri].height = 18

# ════════════════════════════════════════════════════════════════
# HOJA 4: BODEGAS
# ════════════════════════════════════════════════════════════════
ws4 = wb.create_sheet("Bodegas")
ws4.sheet_view.showGridLines = False

ws4.merge_cells("A1:C1")
ws4["A1"] = "INVESAKK SAS · Pedidos Pendientes por Bodega"
ws4["A1"].font = Font(name="Calibri", bold=True, size=13, color=BLANCO)
ws4["A1"].fill = fill(ROJO)
ws4["A1"].alignment = Alignment(horizontal="center", vertical="center")
ws4.row_dimensions[1].height = 28

for ci,(name,ancho) in enumerate([("Sede",22),("Pedidos",12),("NT",10)],1):
    c = ws4.cell(row=2, column=ci, value=name)
    c.font = Font(name="Calibri", bold=True, size=10, color=BLANCO)
    c.fill = fill(ROJO)
    c.alignment = Alignment(horizontal="center",vertical="center")
    c.border = borde_fino()
    ws4.column_dimensions[get_column_letter(ci)].width = ancho
ws4.row_dimensions[2].height = 30

bodegas = [
    ["Barranquilla Centro",42,8],
    ["Barranquilla Norte", 38,5],
    ["Barranquilla Sur",   29,3],
    ["Soledad",            15,2],
    ["Malambo",            11,1],
    ["Bogotá",             9, 0],
    ["Montería",           7, 0],
    ["Valledupar",         6, 1],
    ["Cartagena",          5, 0],
]
for ri, row in enumerate(bodegas, 3):
    for ci, val in enumerate(row, 1):
        c = ws4.cell(row=ri, column=ci, value=val)
        c.font = Font(name="Calibri", size=10)
        c.border = borde_fino()
        c.alignment = Alignment(vertical="center")
        c.fill = fill(ROJO_S if ri%2==0 else BLANCO)
    ws4.row_dimensions[ri].height = 18

# ════════════════════════════════════════════════════════════════
# HOJA 5: GUÍA
# ════════════════════════════════════════════════════════════════
ws5 = wb.create_sheet("Guía de uso")
ws5.sheet_view.showGridLines = False
ws5.column_dimensions["A"].width = 4
ws5.column_dimensions["B"].width = 26
ws5.column_dimensions["C"].width = 50

ws5.merge_cells("B1:C1")
ws5["B1"] = "GUÍA DE CARGA — Gestión Logística INVESAKK SAS"
ws5["B1"].font = Font(name="Calibri", bold=True, size=14, color=BLANCO)
ws5["B1"].fill = fill(AZUL_OSCURO)
ws5["B1"].alignment = Alignment(horizontal="center", vertical="center")
ws5.row_dimensions[1].height = 30

guia = [
    ("", "", ""),
    ("", "¿QUÉ HACER?", ""),
    ("", "1. Hoja Facturas",   "Exporta la consulta SQL a Excel y pega las filas en esta hoja, O usa este archivo completo."),
    ("", "2. Sube el archivo", "Panel Admin → pestaña 'Facturas/Pedidos' → arrastra o selecciona este archivo."),
    ("", "3. Verifica preview","La plataforma muestra las primeras 5 filas. Si está correcto, clic en '✔ Aplicar datos'."),
    ("", "", ""),
    ("", "COLUMNAS CLAVE — HOJA FACTURAS", ""),
    ("", "numero",                   "Número de documento (obligatorio)"),
    ("", "nombre",                   "Nombre del conductor — ej. MARIN (TR) JOSE (obligatorio)"),
    ("", "tipo",                     "NT20 | NT10 | NT6 → Nota traslado (sin bono)  |  FEPA | G1AG → Factura (con bono)"),
    ("", "fec_Factura",              "Fecha de la factura en formato YYYY-MM-DD"),
    ("", "fec_promesa_entrega",      "Fecha prometida de entrega en YYYY-MM-DD"),
    ("", "fec_entrega_logistica",    "Fecha real de entrega. Dejar VACÍA si aún no fue entregada (= Pendiente)"),
    ("", "diferencia_fecha_dias",    "Días de diferencia calculados por SQL. La plataforma los usa directamente para el bono."),
    ("", "", ""),
    ("", "ESQUEMA DE BONOS", ""),
    ("", "≤ 0 días",   "+$ 6.000 por pedido (entrega anticipada o mismo día)"),
    ("", "1 día",      "+$ 6.000 por pedido"),
    ("", "2 días",     "+$ 5.000 por pedido"),
    ("", "3–5 días",   "+$ 4.000 por pedido"),
    ("", "6–10 días",  "−$ 400 por pedido"),
    ("", "11–20 días", "−$ 800 por pedido"),
    ("", "21–30 días", "−$ 2.000 por pedido"),
    ("", "", ""),
    ("", "CONDUCTORES ACTIVOS", "ROA · MARTINEZ · MARIN · MENDOZA · NIÑO"),
    ("", "PLACAS",               "WGX062 · WGX060 · TDU-499 · WGX061 · WGD149"),
    ("", "", ""),
    ("", "NIT empresa",          "802014471-6"),
    ("", "Correo admin",         "mario.pacheco@invesakk.com"),
]

for ri, (_, col_b, col_c) in enumerate(guia, 2):
    cb = ws5.cell(row=ri, column=2, value=col_b)
    cc = ws5.cell(row=ri, column=3, value=col_c)
    ws5.row_dimensions[ri].height = 18
    if col_b in ["¿QUÉ HACER?", "COLUMNAS CLAVE — HOJA FACTURAS", "ESQUEMA DE BONOS", "CONDUCTORES ACTIVOS", "PLACAS", "NIT empresa", "Correo admin"]:
        cb.font = Font(name="Calibri", bold=True, size=10,
                       color=AZUL_OSCURO if "QUÉ" in col_b or "COLUMNAS" in col_b or "ESQUEMA" in col_b else "374151")
        cb.fill = fill(AZUL_CLARO if "QUÉ" in col_b or "COLUMNAS" in col_b or "ESQUEMA" in col_b else GRIS_HEADER)
        cc.fill = fill(AZUL_CLARO if "QUÉ" in col_b or "COLUMNAS" in col_b or "ESQUEMA" in col_b else GRIS_HEADER)
    else:
        cb.font = Font(name="Calibri", size=10, color="374151")
        cc.font = Font(name="Calibri", size=10, color="6B7280")
    cb.alignment = Alignment(vertical="center")
    cc.alignment = Alignment(vertical="center", wrap_text=True)

wb.save(OUT)
print(f"Archivo generado: {OUT}")
print(f"Filas de ejemplo en Facturas: {len(sample_rows)}")
